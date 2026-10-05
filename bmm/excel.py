"""Excel export in the class-workbook layout (spec: reference/bmm_export.py).

Days run down the rows and paths across the columns. The random draws (normal shocks and naive
bootstrap rows) are stored as values, taken from the app's own Simulation, so the workbook and
the app show the same 1,000 paths; everything else is a live formula. Only legacy Excel functions
are used (STDEV, VAR, AVERAGE, SLOPE, PERCENTILE, NORMSDIST, NORMSINV, COUNTIF), so the file also
recalculates in LibreOffice.

Changes from the reference spec, all in rows the spec leaves empty or as extra sheets:
    Inputs!B18   EWMA volatility (lambda in C18), from an EWMA variance column on Prices
    Inputs!B17   "Volatility used" points at the lookback, 30-day or EWMA row the app used
    Inputs!B31   drift limit; FINAL DRIFT (B28) is clamped to +/- B31, as in the app
    Prices!E     Nifty return left blank when a Nifty price is missing (SLOPE skips it)
    RandDemo     20 paths driven by live NORMSINV(RAND()) shocks, for exploring in Excel
"""
from __future__ import annotations

import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from bmm.simulate import Simulation
from bmm.trend import MOMENTUM_DAYS, Trend

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BOLD = Font(name=FONT, bold=True)
PLAIN = Font(name=FONT)
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, color="666666")
HEAD = PatternFill("solid", fgColor="D9E1F2")
KEY = PatternFill("solid", fgColor="FFFF00")
LINE = Border(bottom=Side(style="thin", color="999999"))
INR = '"₹"#,##0.00'
PCT = "0.00%"
RAND_PATHS = 20

DISCLAIMER = ("Educational model of stock price movement only. Not investment advice or a "
              "forecast. Brownian motion assumes constant drift and volatility, normally "
              "distributed log returns and no jumps. The drift is a scenario built from the "
              "inputs shown, not a prediction.")


def _hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font, c.fill, c.border = BOLD, HEAD, LINE


def _put(ws, row, label, value, fmt=None, note="", inp=False, key=False, bold=False):
    ws.cell(row, 1, label).font = BOLD if bold else PLAIN
    c = ws.cell(row, 2, value)
    c.font = PLAIN
    if fmt:
        c.number_format = fmt
    if inp:
        c.font = BLUE
    if bold:
        c.font = Font(name=FONT, bold=True, color="0000FF" if inp else "000000")
    if key:
        c.fill = KEY
    if note:
        ws.cell(row, 5 if ws.title == "Inputs" else 3, note).font = NOTE


def _arial_everywhere(wb: Workbook) -> None:
    """Cells written without an explicit font fall back to the Normal style's font."""
    wb._fonts[0] = Font(name=FONT, size=11)  # noqa: SLF001 - openpyxl keeps no public setter


def filename(ticker: str, horizon: int, as_of: pd.Timestamp) -> str:
    safe = "".join(ch if ch.isalnum() else "_" for ch in ticker)
    return f"BMM_{safe}_{horizon}d_{as_of:%Y-%m-%d}.xlsx"


def build_workbook(*, ticker: str, company: str, prices: pd.Series, index: pd.Series | None,
                   trend: Trend, sim: Simulation, news_note: str = "",
                   data_note: str = "") -> bytes:
    """prices / index: adjusted closes (DatetimeIndex, ascending), the same series the app used
    for `trend` and `sim`."""
    a = trend.assumptions
    prices = prices.dropna()
    if not np.isclose(float(prices.iloc[-1]), sim.spot):
        raise ValueError("sim.spot must be the last price in `prices`")
    idx = index.reindex(prices.index) if index is not None else None
    horizon, n_paths = sim.horizon, sim.paths.shape[1]
    n_obs = len(prices)
    first_ret, last = 3, n_obs + 1           # Prices rows: header 1, data 2..last
    boot_rows = sim.boot_index + first_ret   # Prices row numbers of the drawn returns

    wb = Workbook()
    _arial_everywhere(wb)
    wb.calculation.fullCalcOnLoad = True  # no cached values: Excel computes everything on open

    # ---------------------------------------------------------------- ReadMe
    rd = wb.active
    rd.title = "ReadMe"
    rd["A1"] = f"Brownian motion model of stock price: {company} ({ticker})"
    rd["A1"].font = TITLE
    info = [
        ("Purpose", f"Simulates how the stock price can move over the next {horizon} trading "
                    "days, treating price like a Brownian particle: a drift (trend) plus random "
                    "kicks."),
        ("Physics link", "Drift = steady force on the particle. Volatility x random shock = "
                         "molecular kicks. The spread of outcomes grows with the square root of "
                         "time, like diffusion."),
        ("Equation (Exact)", "S(t+1) = S(t) x EXP((mu - sigma^2/2) x dt + sigma x SQRT(dt) x "
                             "eps),  eps ~ N(0,1)"),
        ("Equation (Euler)", "S(t+1) = S(t) x (1 + mu x dt + sigma x SQRT(dt) x eps)   "
                             "[class workbook method]"),
        ("Naive method", "Class Sheet 1 logic: each day applies a randomly picked historical "
                         "daily log return."),
        ("How to use", "Edit blue cells on 'Inputs'. Yellow = key assumptions. Black = formulas."),
        ("How to verify", "See 'Summary' > Validation: simulated mean, median, P(down) and "
                          "spread vs closed-form theory."),
        ("Data", data_note or "Adjusted closes for the stock and the Nifty 50."),
        ("News", news_note or "News sentiment score entered on Inputs."),
        ("Matches the app", f"Shocks and naive draws are the app's own (seed {sim.seed}), so "
                            "Paths, Naive, Bands and Summary show the same numbers as the app."),
        ("", ""),
        ("Inputs", "Spot, volatility (lookback, 30-day, EWMA), drift components, weights, news "
                   "adjustment, drift limit, final drift"),
        ("Prices", "Lookback adjusted closes for the stock and Nifty 50, daily log returns, "
                   "EWMA variance"),
        ("Shocks", f"{horizon} x {n_paths:,} standard normal draws (seed {sim.seed}), stored as "
                   "values"),
        ("Paths", f"{n_paths:,} Brownian (GBM) paths, days down rows, paths across columns"),
        ("NaiveDraws / Naive", "Random historical-return row numbers and the resulting bootstrap "
                               "paths"),
        ("Bands", "Per-day percentiles, mean and theoretical mean, with charts"),
        ("RandDemo", f"{RAND_PATHS} paths with live RAND() shocks: press F9 to redraw. These do "
                     "not match the app; they show the randomness."),
        ("Summary", "Outcome statistics, probabilities, validation vs theory, technical context"),
    ]
    for i, (k, v) in enumerate(info, start=3):
        rd.cell(i, 1, k).font = BOLD
        c = rd.cell(i, 2, v)
        c.font, c.alignment = PLAIN, Alignment(wrap_text=True, vertical="top")
    disc_row = 3 + len(info) + 1
    rd.cell(disc_row, 1, "Disclaimer").font = BOLD
    c = rd.cell(disc_row, 2, DISCLAIMER)
    c.font, c.alignment = PLAIN, Alignment(wrap_text=True, vertical="top")
    rd.column_dimensions["A"].width = 20
    rd.column_dimensions["B"].width = 110

    # ---------------------------------------------------------------- Prices
    wp = wb.create_sheet("Prices")
    _hdr(wp, 1, ["Date", f"{ticker} adj. close (₹)", "Nifty 50 adj. close",
                 f"{ticker} log return", "Nifty log return", "EWMA variance"])
    seed_row = first_ret + 29                # EWMA seeded with the first 30 returns
    for i, (dt, px) in enumerate(prices.items(), start=2):
        wp.cell(i, 1, pd.Timestamp(dt).to_pydatetime()).number_format = "dd-mmm-yyyy"
        c = wp.cell(i, 2, float(px))
        c.font, c.number_format = BLUE, "#,##0.00"
        if idx is not None and pd.notna(idx.iloc[i - 2]):
            c = wp.cell(i, 3, float(idx.iloc[i - 2]))
            c.font, c.number_format = BLUE, "#,##0.00"
        if i >= first_ret:
            wp.cell(i, 4, f"=LN(B{i}/B{i-1})").number_format = "0.0000%"
            if idx is not None:
                wp.cell(i, 5, f'=IF(AND(ISNUMBER(C{i}),ISNUMBER(C{i-1})),LN(C{i}/C{i-1}),"")'
                        ).number_format = "0.0000%"
        if i == seed_row:
            wp.cell(i, 6, f"=VAR(D{first_ret}:D{seed_row})").number_format = "0.000000"
        elif i > seed_row:
            wp.cell(i, 6, f"=Inputs!$C$18*F{i-1}+(1-Inputs!$C$18)*D{i}^2"
                    ).number_format = "0.000000"
    wp.freeze_panes = "A2"
    for col, w in zip("ABCDEF", (13, 18, 17, 16, 15, 15), strict=True):
        wp.column_dimensions[col].width = w
    D = f"Prices!$D${first_ret}:$D${last}"
    E = f"Prices!$E${first_ret}:$E${last}"

    # ---------------------------------------------------------------- Inputs
    wi = wb.create_sheet("Inputs", 1)
    wi["A1"] = "Inputs: Brownian motion of stock price"
    wi["A1"].font = TITLE
    _hdr(wi, 3, ["Parameter", "Value", "", "", "Note"])
    _put(wi, 4, "Ticker", ticker)
    _put(wi, 5, "Company", company)
    _put(wi, 6, "Last price date", f"=Prices!A{last}", "dd-mmm-yyyy")
    _put(wi, 7, "Spot S0 (₹)", f"=Prices!B{last}", INR, "Latest adjusted close in the data",
         bold=True)
    _put(wi, 8, "Trading days per year", a.days_per_year, "0", "Class convention", inp=True)
    _put(wi, 9, "dt (1 trading day)", "=1/B8", "0.000000")
    _put(wi, 10, "Horizon (trading days)", horizon, "0",
         "Fixed at build: sets the number of path rows")
    _put(wi, 11, "T (years)", "=B10/B8", "0.0000")
    _put(wi, 12, "Simulation method", sim.method, None,
         "Exact (lognormal) or Euler (class method)", inp=True, key=True)
    dv = DataValidation(type="list", formula1='"Exact,Euler"', allow_blank=False)
    wi.add_data_validation(dv)
    dv.add("B12")

    _hdr(wi, 14, ["Volatility (sigma, annual)", "Value", "", "", "Note"])
    _put(wi, 15, "Lookback volatility", f"=STDEV({D})*SQRT(B8)", PCT,
         "Stdev of daily log returns x SQRT(255)")
    recent = f"=STDEV(Prices!D{last-29}:D{last})*SQRT(B8)" if n_obs > 30 else '=""'
    _put(wi, 16, "30-day volatility", recent, PCT, "Recent regime")
    used_row = {"lookback": 15, "30d": 16, "ewma": 18}[trend.sigma_method]
    _put(wi, 17, "Volatility used", f"=B{used_row}", PCT,
         "Points at B15 (lookback), B16 (30-day) or B18 (EWMA)", key=True, bold=True)
    ewma = f"=SQRT(Prices!F{last}*B8)" if last > seed_row else '=""'
    _put(wi, 18, "EWMA volatility (lambda in C18)", ewma, PCT,
         "RiskMetrics: var = lambda x var(prev) + (1 - lambda) x r^2, seeded with the first "
         "30 returns")
    c = wi.cell(18, 3, a.ewma_lambda)
    c.font, c.number_format = BLUE, "0.00"

    _hdr(wi, 19, ["Drift component (annual mu)", "Annual mu", "Weight", "Contribution",
                  "How it is computed"])
    comp_rows = {"historical": 20, "momentum_3m": 21, "momentum_6m": 22, "capm": 23}
    formulas = {
        "historical": (f"=AVERAGE({D})*B8+B17^2/2",
                       "Mean daily log return x 255, + sigma^2/2 to convert log drift to mu"),
        "capm": ("=B33+B34*B35", "r + beta x equity risk premium"),
    }
    for key, days in MOMENTUM_DAYS.items():
        label = "3-month" if days == 63 else "6-month"
        if last - days >= 2:
            formulas[key] = (f"=LN(B7/Prices!B{last-days})*B8/{days}+B17^2/2",
                             f"Last {days} trading days, annualised, + sigma^2/2")
        else:
            formulas[key] = ("n/a", f"Not enough history for {label} momentum: weight set to 0")
    labels = {"historical": "Historical drift", "momentum_3m": "3-month momentum",
              "momentum_6m": "6-month momentum", "capm": "CAPM expected return"}
    for key, rr in comp_rows.items():
        f, note = formulas[key]
        wi.cell(rr, 1, labels[key]).font = PLAIN
        wi.cell(rr, 2, f).number_format = PCT
        weight = 0.0 if f == "n/a" else float(a.weights.get(key, 0))
        c = wi.cell(rr, 3, weight)
        c.font, c.number_format = BLUE, "0%"
        wi.cell(rr, 4, f"=IF(ISNUMBER(B{rr}),B{rr}*C{rr}/SUM($C$20:$C$23),0)"
                ).number_format = PCT
        wi.cell(rr, 5, note).font = NOTE
    _put(wi, 24, "Price-trend mu (weighted)", "=SUMPRODUCT(B20:B23,C20:C23)/SUM(C20:C23)", PCT,
         bold=True)
    wi["C24"] = "=SUM(C20:C23)"
    wi["C24"].number_format = "0%"
    _put(wi, 25, "News sentiment score (-1 to +1)", float(trend.news_score), "0.00",
         "From recent headlines: -1 very bearish, +1 very bullish", inp=True, key=True)
    _put(wi, 26, "News adjustment cap", float(a.news_cap), PCT, "Max drift change from news",
         inp=True)
    _put(wi, 27, "News adjustment", "=B25*B26", PCT)
    _put(wi, 28, "FINAL DRIFT mu (annual)", "=MAX(-B31,MIN(B31,B24+B27))", PCT,
         "Drives every Brownian path; clamped to +/- B31", key=True, bold=True)
    _put(wi, 29, "Trend label", '=IF(B28<-B30,"Bearish",IF(B28>B30,"Bullish","Neutral"))',
         None, bold=True)
    _put(wi, 30, "Neutral band (+/-)", float(a.neutral_band), PCT,
         "Drift within this band = Neutral", inp=True)
    _put(wi, 31, "Drift limit (+/-)", float(a.mu_limit), PCT,
         "Final drift is clamped to this range", inp=True)

    _hdr(wi, 32, ["CAPM inputs", "Value", "", "", "Note"])
    _put(wi, 33, "Risk-free rate r", float(a.risk_free), PCT, a.risk_free_note, inp=True,
         key=True)
    if idx is not None and idx.notna().sum() > 30:
        _put(wi, 34, "Beta vs Nifty 50", f"=SLOPE({D},{E})", "0.000",
             "Regression of daily log returns")
    else:
        _put(wi, 34, "Beta vs Nifty 50", 1.0, "0.000",
             "Nifty 50 data unavailable, so beta is set to 1.0", inp=True)
    _put(wi, 35, "Equity risk premium", float(a.erp), PCT, "Assumption for India", inp=True,
         key=True)

    _hdr(wi, 37, ["Drift vs noise over the horizon", "Value", "", "", "Note"])
    _put(wi, 38, "Log drift per year (mu - sigma^2/2)", "=B28-B17^2/2", PCT)
    _put(wi, 39, "Drift over horizon", "=B38*B11", PCT, "Expected log move from the trend")
    _put(wi, 40, "Noise over horizon (sigma x SQRT(T))", "=B17*SQRT(B11)", PCT,
         "1 standard deviation of log move")
    _put(wi, 41, "Drift / noise ratio", "=IF(B40=0,0,ABS(B39)/B40)", "0.00",
         "Below 1: randomness dominates the trend")

    _hdr(wi, 43, ["Per-step terms used by the paths", "Value", "", "", "Note"])
    _put(wi, 44, "(mu - sigma^2/2) x dt", "=(B28-B17^2/2)*B9", "0.000000",
         "Exact method drift step")
    _put(wi, 45, "sigma x SQRT(dt)", "=B17*SQRT(B9)", "0.000000", "Shock scale per step")
    _put(wi, 46, "mu x dt", "=B28*B9", "0.000000", "Euler method drift step")
    _put(wi, 47, "Random seed", int(sim.seed), "0",
         "Shocks and naive draws generated with this seed (same as the app)")
    for col, w in zip("ABCDE", (36, 16, 10, 14, 70), strict=True):
        wi.column_dimensions[col].width = w

    # ---------------------------------------------------------------- Shocks, Paths, Naive
    ws = wb.create_sheet("Shocks")
    wpa = wb.create_sheet("Paths")
    wnd = wb.create_sheet("NaiveDraws")
    wnv = wb.create_sheet("Naive")
    for sh in (ws, wpa, wnd, wnv):
        sh.cell(1, 1, "Day").font = BOLD
        sh.freeze_panes = "B2"
    lastc = L(n_paths + 1)
    for p in range(n_paths):
        col = p + 2
        for sh in (ws, wpa, wnd, wnv):
            sh.cell(1, col, f"Path {p + 1}").font = BOLD
        wpa.cell(2, col, "=Inputs!$B$7").number_format = "#,##0.00"
        wnv.cell(2, col, "=Inputs!$B$7").number_format = "#,##0.00"
    for d in range(0, horizon + 1):
        for sh in (ws, wpa, wnd, wnv):
            sh.cell(d + 2, 1, d)
    for d in range(1, horizon + 1):
        rr = d + 2
        for p in range(n_paths):
            col = p + 2
            cl = L(col)
            ws.cell(rr, col, float(sim.shocks[d - 1, p])).number_format = "0.0000"
            wpa.cell(rr, col,
                     f'=IF(Inputs!$B$12="Euler",{cl}{rr-1}*(1+Inputs!$B$46+Inputs!$B$45*'
                     f'Shocks!{cl}{rr}),{cl}{rr-1}*EXP(Inputs!$B$44+Inputs!$B$45*Shocks!{cl}{rr}))'
                     ).number_format = "#,##0.00"
            wnd.cell(rr, col, int(boot_rows[d - 1, p]))
            wnv.cell(rr, col, f"={cl}{rr-1}*EXP(INDEX(Prices!$D:$D,NaiveDraws!{cl}{rr}))"
                     ).number_format = "#,##0.00"
    end = horizon + 2                      # terminal row
    lnr = end + 1
    wpa.cell(lnr, 1, "ln(S_T/S0)").font = BOLD
    for p in range(n_paths):
        cl = L(p + 2)
        wpa.cell(lnr, p + 2, f"=LN({cl}{end}/{cl}2)").number_format = "0.0000"
    wnd.cell(end + 2, 1, "Each cell = row number on 'Prices' whose daily log return is applied "
                         "that day").font = NOTE
    T_SCI = f"Paths!$B${end}:${lastc}${end}"
    T_NAI = f"Naive!$B${end}:${lastc}${end}"

    # ---------------------------------------------------------------- Bands
    wbd = wb.create_sheet("Bands")
    _hdr(wbd, 1, ["Day", "5th pct", "25th pct", "Median", "75th pct", "95th pct",
                  "Simulated mean", "Theory mean S0*e^(mu t)", "Naive median"])
    for d in range(0, horizon + 1):
        rr = d + 2
        row_rng = f"Paths!$B${rr}:${lastc}${rr}"
        wbd.cell(rr, 1, d)
        for j, pc in enumerate((0.05, 0.25, 0.5, 0.75, 0.95), start=2):
            wbd.cell(rr, j, f"=PERCENTILE({row_rng},{pc})").number_format = "#,##0.00"
        wbd.cell(rr, 7, f"=AVERAGE({row_rng})").number_format = "#,##0.00"
        wbd.cell(rr, 8, f"=Inputs!$B$7*EXP(Inputs!$B$28*A{rr}*Inputs!$B$9)"
                 ).number_format = "#,##0.00"
        wbd.cell(rr, 9, f"=PERCENTILE(Naive!$B${rr}:${lastc}${rr},0.5)"
                 ).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHI", (6, 11, 11, 11, 11, 11, 14, 20, 13), strict=True):
        wbd.column_dimensions[col].width = w

    ch = LineChart()
    ch.title = "Brownian cone: percentile bands"
    ch.y_axis.title = "Price (₹)"
    ch.x_axis.title = "Trading day"
    ch.add_data(Reference(wbd, min_col=2, max_col=6, min_row=1, max_row=end),
                titles_from_data=True)
    ch.add_data(Reference(wbd, min_col=9, max_col=9, min_row=1, max_row=end),
                titles_from_data=True)
    ch.set_categories(Reference(wbd, min_col=1, min_row=2, max_row=end))
    ch.height, ch.width = 10, 20
    wbd.add_chart(ch, "K2")

    sp = LineChart()
    sp.title = "First 20 Brownian paths"
    sp.y_axis.title = "Price (₹)"
    sp.x_axis.title = "Trading day"
    sp.add_data(Reference(wpa, min_col=2, max_col=min(21, n_paths + 1), min_row=1, max_row=end),
                titles_from_data=True)
    sp.set_categories(Reference(wpa, min_col=1, min_row=2, max_row=end))
    sp.legend = None
    sp.height, sp.width = 10, 20
    wbd.add_chart(sp, "K24")

    # ---------------------------------------------------------------- RandDemo
    rdm = wb.create_sheet("RandDemo")
    rdm["A1"] = "Day"
    rdm["A1"].font = BOLD
    rdm.freeze_panes = "B2"
    for p in range(RAND_PATHS):
        col = p + 2
        rdm.cell(1, col, f"Path {p + 1}").font = BOLD
        rdm.cell(2, col, "=Inputs!$B$7").number_format = "#,##0.00"
    for d in range(0, horizon + 1):
        rdm.cell(d + 2, 1, d)
    for d in range(1, horizon + 1):
        rr = d + 2
        for p in range(RAND_PATHS):
            cl = L(p + 2)
            rdm.cell(rr, p + 2,
                     f'=IF(Inputs!$B$12="Euler",{cl}{rr-1}*(1+Inputs!$B$46+Inputs!$B$45*'
                     f'NORMSINV(RAND())),{cl}{rr-1}*EXP(Inputs!$B$44+Inputs!$B$45*'
                     f'NORMSINV(RAND())))').number_format = "#,##0.00"
    rl = L(RAND_PATHS + 1)
    rdm.cell(end + 2, 1, "Median of the 20 live paths on the last day").font = BOLD
    rdm.cell(end + 2, 6, f"=PERCENTILE(B{end}:{rl}{end},0.5)").number_format = INR
    rdm.cell(end + 3, 1, "Live RAND() shocks: every recalculation (F9) draws new paths. They do "
                         "not match the app or the Paths sheet, which use stored draws."
             ).font = NOTE
    demo = LineChart()
    demo.title = f"{RAND_PATHS} live random paths (press F9 to redraw)"
    demo.y_axis.title = "Price (₹)"
    demo.x_axis.title = "Trading day"
    demo.add_data(Reference(rdm, min_col=2, max_col=RAND_PATHS + 1, min_row=1, max_row=end),
                  titles_from_data=True)
    demo.set_categories(Reference(rdm, min_col=1, min_row=2, max_row=end))
    demo.legend = None
    demo.height, demo.width = 10, 20
    rdm.add_chart(demo, f"{L(RAND_PATHS + 3)}2")

    # ---------------------------------------------------------------- Summary
    su = wb.create_sheet("Summary", 2)
    su["A1"] = f"Summary: {company} over the next {horizon} trading days"
    su["A1"].font = TITLE
    su["A2"] = ('="Trend: "&Inputs!B29&"   |   Final drift mu: "&TEXT(Inputs!B28,"0.0%")&'
                '"   |   Volatility: "&TEXT(Inputs!B17,"0.0%")&"   |   Method: "&Inputs!B12')
    su["A2"].font = BOLD
    _hdr(su, 4, ["Outcome at horizon", "Brownian (scientific)", "Naive bootstrap"])
    rows = [
        ("Spot S0", "=Inputs!B7", "=Inputs!B7", INR),
        ("Mean price", f"=AVERAGE({T_SCI})", f"=AVERAGE({T_NAI})", INR),
        ("Median price", f"=PERCENTILE({T_SCI},0.5)", f"=PERCENTILE({T_NAI},0.5)", INR),
        ("5th percentile", f"=PERCENTILE({T_SCI},0.05)", f"=PERCENTILE({T_NAI},0.05)", INR),
        ("95th percentile", f"=PERCENTILE({T_SCI},0.95)", f"=PERCENTILE({T_NAI},0.95)", INR),
        ("Expected return (mean / S0 - 1)", "=B6/B5-1", "=C6/C5-1", PCT),
        ("90% range: low vs spot", "=B8/B5-1", "=C8/C5-1", PCT),
        ("90% range: high vs spot", "=B9/B5-1", "=C9/C5-1", PCT),
        ("P(price ends below spot)", f'=COUNTIF({T_SCI},"<"&B5)/COUNT({T_SCI})',
         f'=COUNTIF({T_NAI},"<"&C5)/COUNT({T_NAI})', PCT),
        ("P(fall more than 10%)", f'=COUNTIF({T_SCI},"<"&B5*0.9)/COUNT({T_SCI})',
         f'=COUNTIF({T_NAI},"<"&C5*0.9)/COUNT({T_NAI})', PCT),
        ("P(rise more than 10%)", f'=COUNTIF({T_SCI},">"&B5*1.1)/COUNT({T_SCI})',
         f'=COUNTIF({T_NAI},">"&C5*1.1)/COUNT({T_NAI})', PCT),
        ("Paths", f"=COUNT({T_SCI})", f"=COUNT({T_NAI})", "#,##0"),
    ]
    for i, (lab, x, y, fmt) in enumerate(rows, start=5):
        su.cell(i, 1, lab).font = PLAIN
        su.cell(i, 2, x).number_format = fmt
        su.cell(i, 3, y).number_format = fmt
    su["A17"] = ("Naive bootstrap replays historical daily returns, so it carries the lookback's "
                 "own drift, not the trend inputs.")
    su["A17"].font = NOTE

    _hdr(su, 19, ["Validation vs closed-form theory", "Theory", "Simulated", "Gap"])
    val = [
        ("Mean price  S0 x e^(mu T)", "=Inputs!B7*EXP(Inputs!B28*Inputs!B11)", "=B6"),
        ("Median price  S0 x e^((mu - sigma^2/2) T)", "=Inputs!B7*EXP(Inputs!B38*Inputs!B11)",
         "=B7"),
        ("P(ends below spot)  N(-(mu - sigma^2/2)T / sigma SQRT(T))",
         "=IF(Inputs!B40=0,IF(Inputs!B39<0,1,0),NORMSDIST(-Inputs!B39/Inputs!B40))", "=B13"),
        ("Stdev of ln(S_T/S0)  sigma x SQRT(T)", "=Inputs!B40",
         f"=STDEV(Paths!$B${lnr}:${lastc}${lnr})"),
        ("Mean of ln(S_T/S0)  (mu - sigma^2/2) T", "=Inputs!B39",
         f"=AVERAGE(Paths!$B${lnr}:${lastc}${lnr})"),
    ]
    for i, (lab, th, simf) in enumerate(val, start=20):
        su.cell(i, 1, lab).font = PLAIN
        fmt = INR if i in (20, 21) else PCT
        su.cell(i, 2, th).number_format = fmt
        su.cell(i, 3, simf).number_format = fmt
        su.cell(i, 4, f"=IF(B{i}=0,0,C{i}/B{i}-1)").number_format = "0.0%"
    su["A25"] = ("Gaps of a few percent are sampling noise from 1,000 paths. Theory rows assume "
                 "the Exact method; Euler shifts them slightly.")
    su["A25"].font = NOTE

    _hdr(su, 27, ["Technical context (from Prices)", "Value", "Read"])
    tech = [
        ("52-week high", f"=MAX(Prices!B{max(2, last-254)}:B{last})", None, INR),
        ("Distance from 52-week high", "=Inputs!B7/B28-1", None, PCT),
        ("50-day average", f"=AVERAGE(Prices!B{max(2, last-49)}:B{last})",
         '=IF(Inputs!B7<B30,"Below","Above")', INR),
        ("200-day average", f"=AVERAGE(Prices!B{max(2, last-199)}:B{last})",
         '=IF(Inputs!B7<B31,"Below","Above")', INR),
        ("30-day vs lookback volatility", '=IF(ISNUMBER(Inputs!B16),Inputs!B16/Inputs!B15-1,"")',
         '=IF(B32="","",IF(B32>0.15,"Elevated",IF(B32<-0.15,"Subdued","In line")))', PCT),
    ]
    for i, (lab, f, read, fmt) in enumerate(tech, start=28):
        su.cell(i, 1, lab).font = PLAIN
        su.cell(i, 2, f).number_format = fmt
        if read:
            su.cell(i, 3, read)
    su["A34"] = "Disclaimer"
    su["A34"].font = BOLD
    su["B34"] = DISCLAIMER
    su["B34"].alignment = Alignment(wrap_text=True, vertical="top")
    su.merge_cells("B34:D34")
    su.row_dimensions[34].height = 60
    for col, w in zip("ABCD", (58, 20, 20, 10), strict=True):
        su.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
