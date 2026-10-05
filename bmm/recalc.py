"""Recalculate an .xlsx with LibreOffice and report formula errors.

openpyxl writes formulas without cached values; this opens the file in headless LibreOffice,
runs calculateAll(), saves it in place and then scans every cell for Excel error values.
Used by the Excel tests and `python -m bmm.recalc file.xlsx`.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#NULL!")
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1"
 script:language="StarBasic">
Sub RecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""


def soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def calc_available() -> bool:
    """LibreOffice with its spreadsheet component (Calc) is installed."""
    exe = soffice()
    if exe is None:
        return False
    program = Path(os.path.realpath(exe)).parent
    return any(program.glob("libsclo.*"))  # Calc itself, not the core-only libscnlo


def recalc(path: str | Path, timeout: int = 300) -> dict:
    exe = soffice()
    if exe is None:
        raise FileNotFoundError("LibreOffice (soffice) is not installed")
    path = Path(path).resolve()
    env = {**os.environ, "SAL_USE_VCLPLUGIN": "svp"}
    with tempfile.TemporaryDirectory(prefix="lo_profile_") as profile:
        url = Path(profile).as_uri()
        subprocess.run([exe, f"-env:UserInstallation={url}", "--headless",
                        "--terminate_after_init"], env=env, capture_output=True, timeout=120)
        macro_dir = Path(profile) / "user" / "basic" / "Standard"
        if not macro_dir.exists():
            raise RuntimeError("LibreOffice did not create a user profile")
        (macro_dir / "Module1.xba").write_text(MACRO)
        before = path.stat().st_mtime_ns
        run = subprocess.run(
            [exe, f"-env:UserInstallation={url}", "--headless", "--norestore",
             "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic"
             "&location=application", str(path)],
            env=env, capture_output=True, text=True, timeout=timeout)
        if path.stat().st_mtime_ns == before:
            raise RuntimeError(f"LibreOffice did not save the file: {run.stderr.strip()}")
    return scan(path)


def scan(path: str | Path) -> dict:
    formulas = load_workbook(path, data_only=False, read_only=True)
    values = load_workbook(path, data_only=True, read_only=True)
    total, errors, empty = 0, {}, 0
    for name in formulas.sheetnames:
        for frow, vrow in zip(formulas[name].iter_rows(), values[name].iter_rows(), strict=False):
            for f, v in zip(frow, vrow, strict=False):
                if isinstance(f.value, str) and f.value.startswith("="):
                    total += 1
                    if isinstance(v.value, str) and v.value in ERRORS:
                        errors.setdefault(v.value, []).append(f"{name}!{f.coordinate}")
                    elif v.value is None:
                        empty += 1
    return {"total_formulas": total, "total_errors": sum(map(len, errors.values())),
            "errors": {k: v[:20] for k, v in errors.items()}, "formulas_without_value": empty}


if __name__ == "__main__":
    print(json.dumps(recalc(sys.argv[1]), indent=2))
