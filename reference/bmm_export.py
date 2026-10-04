"""Brownian Motion Model (BMM) workbook for stock price movement.

Layout follows the GIM class workbook: days run down the rows, paths run across the columns.
Random draws are stored as values (reproducible); everything else is a live formula.
Only legacy Excel functions are used (STDEV, AVERAGE, SLOPE, PERCENTILE, NORMSDIST, COUNTIF).
"""
from __future__ import annotations

import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
NOTE = Font(italic=True, color="666666")
HEAD = PatternFill("solid", fgColor="D9E1F2")
KEY = PatternFill("solid", fgColor="FFFF00")
LINE = Border(bottom=Side(style="thin", color="999999"))
INR = '"₹"#,##0.00'
PCT = "0.00%"

DISCLAIMER = ("Educational model of stock price movement only. Not investment advice or a forecast. "
              "Brownian motion assumes constant drift and volatility, normally distributed log returns "
              "and no jumps. The drift is a scenario built from the inputs shown, not a prediction.")


def _hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font, c.fill, c.border = BOLD, HEAD, LINE


def _put(ws, row, label, value, fmt=None, note="", inp=False, key=False, bold=False):
    ws.cell(row, 1, label).font = BOLD if bold else Font()
    c = ws.cell(row, 2, value)
    if fmt:
        c.number_format = fmt
    if inp:
        c.font = BLUE
    if bold:
        c.font = Font(bold=True, color="0000FF" if inp else "000000")
    if key:
        c.fill = KEY
    if note:
        ws.cell(row, 5 if ws.title == "Inputs" else 3, note).font = NOTE


def build_bmm_workbook(*, ticker: str, company: str, prices: pd.DataFrame, horizon: int = 60,
                       n_paths: int = 1000, seed: int = 42, r: float = 0.0534, erp: float = 0.07,
                       news_score: float = 0.0, news_note: str = "", weights=(0.25, 0.25, 0.25, 0.25),
                       method: str = "Exact", data_note: str = "") -> bytes:
    """prices: DataFrame with columns date, close, index_close (ascending, lookback window only)."""
    rng = np.random.default_rng(seed)
    shocks = rng.standard_normal((horizon, n_paths))               # rows = days, cols = paths
    n_obs = len(prices)
    first_ret, last = 3, n_obs + 1                                  # Prices rows: header=1, data 2..last
    boot_rows = rng.integers(first_ret, last + 1, size=(horizon, n_paths))  # rows of historical returns

    wb = Workbook()

    # ---------------------------------------------------------------- ReadMe
    rd = wb.active
    rd.title = "ReadMe"
    rd["A1"] = f"Brownian motion model of stock price: {company} ({ticker})"
    rd["A1"].font = TITLE
    info = [
        ("Purpose", "Simulates how the stock price can move over the next "
                    f"{horizon} trading days, treating price like a Brownian particle: a drift (trend) plus random kicks."),
        ("Physics link", "Drift = steady force on the particle. Volatility x random shock = molecular kicks. "
                         "The spread of outcomes grows with the square root of time, like diffusion."),
        ("Equation (Exact)", "S(t+1) = S(t) x EXP((mu - sigma^2/2) x dt + sigma x SQRT(dt) x eps),  eps ~ N(0,1)"),
        ("Equation (Euler)", "S(t+1) = S(t) x (1 + mu x dt + sigma x SQRT(dt) x eps)   [class workbook method]"),
        ("Naive method", "Class Sheet 1 logic: each day applies a randomly picked historical daily log return."),
        ("How to use", "Edit blue cells on 'Inputs'. Yellow = key assumptions. Black = formulas."),
        ("How to verify", "See 'Summary' > Validation: simulated mean, median, P(down) and spread vs closed-form theory."),
        ("Data", data_note or "Adjusted daily closes."),
        ("News", news_note or "News sentiment score entered on Inputs."),
        ("", ""),
        ("Inputs", "Spot, volatility, drift components, weights, news adjustment, final drift"),
        ("Prices", "Lookback closes for the stock and Nifty 50, daily log returns"),
        ("Shocks", f"{horizon} x {n_paths:,} standard normal draws (seed {seed}), stored as values"),
        ("Paths", f"{n_paths:,} Brownian (GBM) paths, days down rows, paths across columns"),
        ("NaiveDraws / Naive", "Random historical-return row numbers and the resulting bootstrap paths"),
        ("Bands", "Per-day percentiles, mean and theoretical mean, with charts"),
        ("Summary", "Outcome statistics, probabilities, validation vs theory, technical context"),
        ("", ""),
        ("Disclaimer", DISCLAIMER),
    ]
    for i, (a, b) in enumerate(info, start=3):
        rd.cell(i, 1, a).font = BOLD
        rd.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    rd.column_dimensions["A"].width = 20
    rd.column_dimensions["B"].width = 110

    # ---------------------------------------------------------------- Prices
    wp = wb.create_sheet("Prices")
    _hdr(wp, 1, ["Date", f"{ticker} close (₹)", "Nifty 50 close", f"{ticker} log return", "Nifty log return"])
    for i, row in enumerate(prices.itertuples(index=False), start=2):
        wp.cell(i, 1, pd.Timestamp(row.date).to_pydatetime()).number_format = "dd-mmm-yyyy"
        c = wp.cell(i, 2, float(row.close)); c.font = BLUE; c.number_format = "#,##0.00"
        c = wp.cell(i, 3, float(row.index_close)); c.font = BLUE; c.number_format = "#,##0.00"
        if i >= first_ret:
            wp.cell(i, 4, f"=LN(B{i}/B{i-1})").number_format = "0.0000%"
            wp.cell(i, 5, f"=LN(C{i}/C{i-1})").number_format = "0.0000%"
    wp.freeze_panes = "A2"
    for col, w in zip("ABCDE", (13, 16, 15, 16, 15)):
        wp.column_dimensions[col].width = w
    D = f"Prices!$D${first_ret}:$D${last}"
    E = f"Prices!$E${first_ret}:$E${last}"

    # ---------------------------------------------------------------- Inputs
    wi = wb.create_sheet("Inputs", 1)
    wi["A1"] = "Inputs: Brownian motion of stock price"
    wi["A1"].font = TITLE
    _hdr(wi, 3, ["Parameter", "Value", "", "", "Note"])
    _put(wi, 4, "Ticker", ticker)
    _put(wi, 5, "Company", company)
    _put(wi, 6, "Last price date", f"=Prices!A{last}", "dd-mmm-yyyy")
    _put(wi, 7, "Spot S0 (₹)", f"=Prices!B{last}", INR, "Latest close in the data", bold=True)
    _put(wi, 8, "Trading days per year", 255, "0", "Class convention", inp=True)
    _put(wi, 9, "dt (1 trading day)", "=1/B8", "0.000000")
    _put(wi, 10, "Horizon (trading days)", horizon, "0", "Fixed at build: sets the number of path rows")
    _put(wi, 11, "T (years)", "=B10/B8", "0.0000")
    _put(wi, 12, "Simulation method", method, None, "Exact (lognormal) or Euler (class method)", inp=True, key=True)
    dv = DataValidation(type="list", formula1='"Exact,Euler"', allow_blank=False)
    wi.add_data_validation(dv)
    dv.add("B12")

    _hdr(wi, 14, ["Volatility (sigma, annual)", "Value", "", "", "Note"])
    _put(wi, 15, "Lookback volatility", f"=STDEV({D})*SQRT(B8)", PCT, "Stdev of daily log returns x SQRT(255)")
    _put(wi, 16, "30-day volatility", f"=STDEV(Prices!D{last-29}:D{last})*SQRT(B8)", PCT, "Recent regime")
    _put(wi, 17, "Volatility used", "=B15", PCT, "Overwrite with =B16 to use the recent regime", key=True, bold=True)

    _hdr(wi, 19, ["Drift component (annual mu)", "Annual mu", "Weight", "Contribution", "How it is computed"])
    comps = [
        (20, "Historical drift", f"=AVERAGE({D})*B8+B17^2/2", weights[0],
         "Mean daily log return x 255, + sigma^2/2 to convert log drift to mu"),
        (21, "3-month momentum", f"=LN(B7/Prices!B{last-63})*B8/63+B17^2/2", weights[1],
         "Last 63 trading days, annualised, + sigma^2/2"),
        (22, "6-month momentum", f"=LN(B7/Prices!B{last-126})*B8/126+B17^2/2", weights[2],
         "Last 126 trading days, annualised, + sigma^2/2"),
        (23, "CAPM expected return", "=B33+B34*B35", weights[3], "r + beta x equity risk premium"),
    ]
    for rr, lab, f, w, note in comps:
        wi.cell(rr, 1, lab)
        wi.cell(rr, 2, f).number_format = PCT
        c = wi.cell(rr, 3, w); c.font = BLUE; c.number_format = "0%"
        wi.cell(rr, 4, f"=B{rr}*C{rr}/SUM($C$20:$C$23)").number_format = PCT
        wi.cell(rr, 5, note).font = NOTE
    _put(wi, 24, "Price-trend mu (weighted)", "=SUMPRODUCT(B20:B23,C20:C23)/SUM(C20:C23)", PCT, bold=True)
    wi["C24"] = "=SUM(C20:C23)"; wi["C24"].number_format = "0%"
    _put(wi, 25, "News sentiment score (-1 to +1)", float(news_score), "0.00",
         "From recent headlines: -1 very bearish, +1 very bullish", inp=True, key=True)
    _put(wi, 26, "News adjustment cap", 0.10, PCT, "Max drift change from news", inp=True)
    _put(wi, 27, "News adjustment", "=B25*B26", PCT)
    _put(wi, 28, "FINAL DRIFT mu (annual)", "=B24+B27", PCT, "Drives every Brownian path", key=True, bold=True)
    _put(wi, 29, "Trend label", '=IF(B28<-B30,"Bearish",IF(B28>B30,"Bullish","Neutral"))', None, bold=True)
    _put(wi, 30, "Neutral band (+/-)", 0.05, PCT, "Drift within this band = Neutral", inp=True)

    _hdr(wi, 32, ["CAPM inputs", "Value", "", "", "Note"])
    _put(wi, 33, "Risk-free rate r", float(r), PCT, "91-day T-bill cut-off yield (RBI)", inp=True, key=True)
    _put(wi, 34, "Beta vs Nifty 50", f"=SLOPE({D},{E})", "0.000", "Regression of daily log returns")
    _put(wi, 35, "Equity risk premium", float(erp), PCT, "Assumption for India", inp=True, key=True)

    _hdr(wi, 37, ["Drift vs noise over the horizon", "Value", "", "", "Note"])
    _put(wi, 38, "Log drift per year (mu - sigma^2/2)", "=B28-B17^2/2", PCT)
    _put(wi, 39, "Drift over horizon", "=B38*B11", PCT, "Expected log move from the trend")
    _put(wi, 40, "Noise over horizon (sigma x SQRT(T))", "=B17*SQRT(B11)", PCT, "1 standard deviation of log move")
    _put(wi, 41, "Drift / noise ratio", "=ABS(B39)/B40", "0.00", "Below 1: randomness dominates the trend")

    _hdr(wi, 43, ["Per-step terms used by the paths", "Value", "", "", "Note"])
    _put(wi, 44, "(mu - sigma^2/2) x dt", "=(B28-B17^2/2)*B9", "0.000000", "Exact method drift step")
    _put(wi, 45, "sigma x SQRT(dt)", "=B17*SQRT(B9)", "0.000000", "Shock scale per step")
    _put(wi, 46, "mu x dt", "=B28*B9", "0.000000", "Euler method drift step")
    _put(wi, 47, "Random seed", int(seed), "0", "Shocks and naive draws generated with this seed")
    for col, w in zip("ABCDE", (36, 16, 10, 14, 62)):
        wi.column_dimensions[col].width = w

    # ---------------------------------------------------------------- Shocks, Paths
    ws = wb.create_sheet("Shocks")
    wpa = wb.create_sheet("Paths")
    wnd = wb.create_sheet("NaiveDraws")
    wnv = wb.create_sheet("Naive")
    for sh, lab in ((ws, "Day"), (wpa, "Day"), (wnd, "Day"), (wnv, "Day")):
        sh.cell(1, 1, lab).font = BOLD
        sh.freeze_panes = "B2"
    lastc = L(n_paths + 1)
    for p in range(n_paths):
        col = p + 2
        for sh in (ws, wpa, wnd, wnv):
            sh.cell(1, col, f"Path {p + 1}").font = BOLD
        wpa.cell(2, col, "=Inputs!$B$7").number_format = "#,##0.00"
        wnv.cell(2, col, "=Inputs!$B$7").number_format = "#,##0.00"
    for d in range(0, horizon + 1):
        rr = d + 2
        for sh in (ws, wpa, wnd, wnv):
            sh.cell(rr, 1, d)
    for d in range(1, horizon + 1):
        rr = d + 2
        for p in range(n_paths):
            col = p + 2
            cl = L(col)
            ws.cell(rr, col, float(shocks[d - 1, p])).number_format = "0.0000"
            wpa.cell(rr, col,
                     f'=IF(Inputs!$B$12="Euler",{cl}{rr-1}*(1+Inputs!$B$46+Inputs!$B$45*Shocks!{cl}{rr}),'
                     f'{cl}{rr-1}*EXP(Inputs!$B$44+Inputs!$B$45*Shocks!{cl}{rr}))').number_format = "#,##0.00"
            wnd.cell(rr, col, int(boot_rows[d - 1, p]))
            wnv.cell(rr, col, f"={cl}{rr-1}*EXP(INDEX(Prices!$D:$D,NaiveDraws!{cl}{rr}))").number_format = "#,##0.00"
    end = horizon + 2                                   # terminal row
    lnr = end + 1
    wpa.cell(lnr, 1, "ln(S_T/S0)").font = BOLD
    for p in range(n_paths):
        cl = L(p + 2)
        wpa.cell(lnr, p + 2, f"=LN({cl}{end}/{cl}2)").number_format = "0.0000"
    wnd.cell(end + 2, 1, "Each cell = row number on 'Prices' whose daily log return is applied that day").font = NOTE
    T_SCI = f"Paths!$B${end}:${lastc}${end}"
    T_NAI = f"Naive!$B${end}:${lastc}${end}"

    # ---------------------------------------------------------------- Bands
    wbd = wb.create_sheet("Bands")
    _hdr(wbd, 1, ["Day", "5th pct", "25th pct", "Median", "75th pct", "95th pct", "Simulated mean",
                  "Theory mean S0*e^(mu t)", "Naive median"])
    for d in range(0, horizon + 1):
        rr, src = d + 2, d + 2
        row_rng = f"Paths!$B${src}:${lastc}${src}"
        wbd.cell(rr, 1, d)
        for j, pc in enumerate((0.05, 0.25, 0.5, 0.75, 0.95), start=2):
            wbd.cell(rr, j, f"=PERCENTILE({row_rng},{pc})").number_format = "#,##0.00"
        wbd.cell(rr, 7, f"=AVERAGE({row_rng})").number_format = "#,##0.00"
        wbd.cell(rr, 8, f"=Inputs!$B$7*EXP(Inputs!$B$28*A{rr}*Inputs!$B$9)").number_format = "#,##0.00"
        wbd.cell(rr, 9, f"=PERCENTILE(Naive!$B${src}:${lastc}${src},0.5)").number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHI", (6, 11, 11, 11, 11, 11, 14, 20, 13)):
        wbd.column_dimensions[col].width = w

    ch = LineChart()
    ch.title = "Brownian cone: percentile bands"
    ch.y_axis.title = "Price (₹)"
    ch.x_axis.title = "Trading day"
    ch.add_data(Reference(wbd, min_col=2, max_col=6, min_row=1, max_row=end), titles_from_data=True)
    ch.add_data(Reference(wbd, min_col=9, max_col=9, min_row=1, max_row=end), titles_from_data=True)
    ch.set_categories(Reference(wbd, min_col=1, min_row=2, max_row=end))
    ch.height, ch.width = 10, 20
    wbd.add_chart(ch, "K2")

    sp = LineChart()
    sp.title = "First 20 Brownian paths"
    sp.y_axis.title = "Price (₹)"
    sp.x_axis.title = "Trading day"
    sp.add_data(Reference(wpa, min_col=2, max_col=21, min_row=1, max_row=end), titles_from_data=True)
    sp.set_categories(Reference(wpa, min_col=1, min_row=2, max_row=end))
    sp.legend = None
    sp.height, sp.width = 10, 20
    wbd.add_chart(sp, "K24")

    # ---------------------------------------------------------------- Summary
    su = wb.create_sheet("Summary", 2)
    su["A1"] = f"Summary: {company} over the next {horizon} trading days"
    su["A1"].font = TITLE
    su["A2"] = '="Trend: "&Inputs!B29&"   |   Final drift mu: "&TEXT(Inputs!B28,"0.0%")&"   |   Volatility: "&TEXT(Inputs!B17,"0.0%")&"   |   Method: "&Inputs!B12'
    su["A2"].font = BOLD
    _hdr(su, 4, ["Outcome at horizon", "Brownian (scientific)", "Naive bootstrap"])
    rows = [
        ("Spot S0", "=Inputs!B7", "=Inputs!B7", INR),
        ("Mean price", f"=AVERAGE({T_SCI})", f"=AVERAGE({T_NAI})", INR),
        ("Median price", f"=PERCENTILE({T_SCI},0.5)", f"=PERCENTILE({T_NAI},0.5)", INR),
        ("5th percentile", f"=PERCENTILE({T_SCI},0.05)", f"=PERCENTILE({T_NAI},0.05)", INR),
        ("95th percentile", f"=PERCENTILE({T_SCI},0.95)", f"=PERCENTILE({T_NAI},0.95)", INR),
        ("Expected return (mean / S0 - 1)", "=B6/B5-1", "=C6/C5-1", PCT),
        ("90% range: low vs spot", "=B8/B5-1", "=C8/C5-1", PCT),
        ("90% range: high vs spot", "=B9/B5-1", "=C9/C5-1", PCT),
        ("P(price ends below spot)", f'=COUNTIF({T_SCI},"<"&B5)/COUNT({T_SCI})',
         f'=COUNTIF({T_NAI},"<"&C5)/COUNT({T_NAI})', PCT),
        ("P(fall more than 10%)", f'=COUNTIF({T_SCI},"<"&B5*0.9)/COUNT({T_SCI})',
         f'=COUNTIF({T_NAI},"<"&C5*0.9)/COUNT({T_NAI})', PCT),
        ("P(rise more than 10%)", f'=COUNTIF({T_SCI},">"&B5*1.1)/COUNT({T_SCI})',
         f'=COUNTIF({T_NAI},">"&C5*1.1)/COUNT({T_NAI})', PCT),
        ("Paths", f"=COUNT({T_SCI})", f"=COUNT({T_NAI})", "#,##0"),
    ]
    for i, (lab, a, b, fmt) in enumerate(rows, start=5):
        su.cell(i, 1, lab)
        su.cell(i, 2, a).number_format = fmt
        su.cell(i, 3, b).number_format = fmt
    su["A17"] = "Naive bootstrap replays historical daily returns, so it carries the lookback's own drift, not the trend inputs."
    su["A17"].font = NOTE

    _hdr(su, 19, ["Validation vs closed-form theory", "Theory", "Simulated", "Gap"])
    val = [
        ("Mean price  S0 x e^(mu T)", "=Inputs!B7*EXP(Inputs!B28*Inputs!B11)", "=B6"),
        ("Median price  S0 x e^((mu - sigma^2/2) T)", "=Inputs!B7*EXP(Inputs!B38*Inputs!B11)", "=B7"),
        ("P(ends below spot)  N(-(mu - sigma^2/2)T / sigma SQRT(T))", "=NORMSDIST(-Inputs!B39/Inputs!B40)", "=B13"),
        ("Stdev of ln(S_T/S0)  sigma x SQRT(T)", "=Inputs!B40", f"=STDEV(Paths!$B${lnr}:${lastc}${lnr})"),
        ("Mean of ln(S_T/S0)  (mu - sigma^2/2) T", "=Inputs!B39", f"=AVERAGE(Paths!$B${lnr}:${lastc}${lnr})"),
    ]
    for i, (lab, th, sim) in enumerate(val, start=20):
        su.cell(i, 1, lab)
        fmt = INR if i in (20, 21) else PCT
        su.cell(i, 2, th).number_format = fmt
        su.cell(i, 3, sim).number_format = fmt
        su.cell(i, 4, f"=IF(B{i}=0,0,C{i}/B{i}-1)").number_format = "0.0%"
    su["A25"] = ("Gaps of a few percent are sampling noise from 1,000 paths. Theory rows assume the Exact method; "
                 "Euler shifts them slightly.")
    su["A25"].font = NOTE

    _hdr(su, 27, ["Technical context (from Prices)", "Value", "Read"])
    tech = [
        ("52-week high", f"=MAX(Prices!B{max(2, last-254)}:B{last})", None, INR),
        ("Distance from 52-week high", "=Inputs!B7/B28-1", None, PCT),
        ("50-day average", f"=AVERAGE(Prices!B{last-49}:B{last})", '=IF(Inputs!B7<B30,"Below","Above")', INR),
        ("200-day average", f"=AVERAGE(Prices!B{last-199}:B{last})", '=IF(Inputs!B7<B31,"Below","Above")', INR),
        ("30-day vs lookback volatility", "=Inputs!B16/Inputs!B15-1",
         '=IF(B32>0.15,"Elevated",IF(B32<-0.15,"Subdued","In line"))', PCT),
    ]
    for i, (lab, f, read, fmt) in enumerate(tech, start=28):
        su.cell(i, 1, lab)
        su.cell(i, 2, f).number_format = fmt
        if read:
            su.cell(i, 3, read)
    su["A34"] = "Disclaimer"
    su["A34"].font = BOLD
    su["B34"] = DISCLAIMER
    su["B34"].alignment = Alignment(wrap_text=True, vertical="top")
    su.merge_cells("B34:D34")
    su.row_dimensions[34].height = 60
    for col, w in zip("ABCD", (58, 20, 20, 10)):
        su.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
