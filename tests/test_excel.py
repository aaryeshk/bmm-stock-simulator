"""Excel export: layout and formulas (always), recalculated values (needs LibreOffice Calc)."""
import io
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from bmm.excel import build_workbook, filename
from bmm.recalc import calc_available, recalc
from bmm.simulate import simulate, summary
from bmm.trend import Assumptions, compute_trend, log_returns

FIX = Path(__file__).parent / "fixtures"
INPUTS = json.loads((FIX / "class_sample_inputs.json").read_text())
SIM = json.loads((FIX / "class_sample_simulation.json").read_text())
needs_calc = pytest.mark.skipif(not calc_available(), reason="LibreOffice Calc not installed")


def build(stock, index, news=0.0, assumptions=None, horizon=60, seed=42, method="Exact"):
    t = compute_trend(stock, index, news, assumptions)
    sim = simulate(t.spot, t.mu, t.sigma, log_returns(stock).to_numpy(), horizon, 1000, seed,
                   method)
    data = build_workbook(ticker="RELIANCE", company="Reliance Industries Ltd.", prices=stock,
                          index=index, trend=t, sim=sim)
    return t, sim, data


@pytest.fixture(scope="module")
def class_data():
    df = pd.read_csv(FIX / "class_reliance_nifty.csv", index_col="date", parse_dates=["date"])
    return df["close"], df["index_close"]


@pytest.fixture(scope="module")
def class_book(class_data):
    return build(*class_data, news=-0.5)


@pytest.fixture(scope="module")
def class_recalculated(class_book, tmp_path_factory):
    if not calc_available():
        pytest.skip("LibreOffice Calc not installed")
    path = tmp_path_factory.mktemp("xl") / "class.xlsx"
    path.write_bytes(class_book[2])
    report = recalc(path)
    return report, load_workbook(path, data_only=True)


# --- Layout and formulas (no LibreOffice needed) ------------------------------------------------

def test_sheets_follow_the_class_layout(class_book):
    wb = load_workbook(io.BytesIO(class_book[2]))
    assert wb.sheetnames == ["ReadMe", "Inputs", "Summary", "Prices", "Shocks", "Paths",
                             "NaiveDraws", "Naive", "Bands", "RandDemo"]
    assert wb.calculation.fullCalcOnLoad


def test_path_formulas_are_the_reference_formulas(class_book):
    wb = load_workbook(io.BytesIO(class_book[2]))
    assert wb["Paths"]["B3"].value == ('=IF(Inputs!$B$12="Euler",B2*(1+Inputs!$B$46+'
                                       'Inputs!$B$45*Shocks!B3),B2*EXP(Inputs!$B$44+'
                                       'Inputs!$B$45*Shocks!B3))')
    assert wb["Naive"]["ALM62"].value == "=ALM61*EXP(INDEX(Prices!$D:$D,NaiveDraws!ALM62))"
    assert wb["Inputs"]["B15"].value == "=STDEV(Prices!$D$3:$D$497)*SQRT(B8)"
    assert wb["Inputs"]["B34"].value == "=SLOPE(Prices!$D$3:$D$497,Prices!$E$3:$E$497)"
    assert wb["Inputs"]["B28"].value == "=MAX(-B31,MIN(B31,B24+B27))"


def test_stored_draws_are_the_apps(class_book):
    _, sim, data = class_book
    wb = load_workbook(io.BytesIO(data), read_only=True)
    shocks = np.array([r for r in wb["Shocks"].iter_rows(min_row=3, max_row=62, min_col=2,
                                                          max_col=1001, values_only=True)])
    draws = np.array([r for r in wb["NaiveDraws"].iter_rows(min_row=3, max_row=62, min_col=2,
                                                             max_col=1001, values_only=True)])
    # openpyxl writes 16 significant digits: at most one unit in the last place.
    assert np.allclose(shocks, sim.shocks, rtol=1e-15, atol=1e-15)
    assert np.array_equal(draws, sim.boot_index + 3)


def test_inputs_carry_the_app_assumptions(class_data):
    a = Assumptions(risk_free=0.06, erp=0.065, news_cap=0.08, neutral_band=0.03, mu_limit=0.25,
                    sigma_method="ewma",
                    weights={"historical": 0.1, "momentum_3m": 0.2, "momentum_6m": 0.3,
                             "capm": 0.4})
    t, sim, data = build(*class_data, news=0.4, assumptions=a, horizon=90, seed=7,
                         method="Euler")
    wi = load_workbook(io.BytesIO(data))["Inputs"]
    assert [wi[c].value for c in ("B33", "B35", "B26", "B30", "B31", "B25")] == \
        [0.06, 0.065, 0.08, 0.03, 0.25, pytest.approx(0.4)]
    assert [wi[f"C{r}"].value for r in range(20, 24)] == [0.1, 0.2, 0.3, 0.4]
    assert wi["B17"].value == "=B18" and wi["B12"].value == "Euler"
    assert wi["B10"].value == 90 and wi["B47"].value == 7


def test_rand_demo_sheet_uses_live_random_shocks(class_book):
    wb = load_workbook(io.BytesIO(class_book[2]))
    demo = wb["RandDemo"]
    assert demo.max_column == 21 and "NORMSINV(RAND())" in demo["B3"].value
    assert demo._charts


def test_arial_and_blue_inputs(class_book):
    wb = load_workbook(io.BytesIO(class_book[2]))
    assert wb["Inputs"]["A1"].font.name == "Arial"
    assert wb["Inputs"]["B33"].font.color.rgb.endswith("0000FF")   # risk-free rate is an input


def test_filename():
    assert filename("M&M", 60, pd.Timestamp("2026-10-01")) == "BMM_M_M_60d_2026-10-01.xlsx"


# --- Recalculated in LibreOffice ----------------------------------------------------------------

@needs_calc
def test_zero_recalc_errors(class_recalculated):
    report, _ = class_recalculated
    assert report["total_errors"] == 0, report["errors"]
    assert report["total_formulas"] > 120_000


@needs_calc
def test_recalculated_inputs_match_the_verified_sample(class_recalculated):
    _, wb = class_recalculated
    wi, su = wb["Inputs"], wb["Summary"]
    cells = {"sigma_lookback": wi["B15"], "sigma_30d": wi["B16"], "historical": wi["B20"],
             "momentum_3m": wi["B21"], "momentum_6m": wi["B22"], "capm": wi["B23"],
             "price_mu": wi["B24"], "mu": wi["B28"], "beta": wi["B34"], "noise": wi["B40"],
             "high_52w": su["B28"], "ma_200": su["B31"]}
    for key, cell in cells.items():
        assert cell.value == pytest.approx(INPUTS[key], rel=1e-12), key
    assert wi["B29"].value == "Bearish"


@needs_calc
def test_recalculated_summary_matches_the_verified_sample(class_recalculated):
    _, wb = class_recalculated
    su = wb["Summary"]
    keys = ["spot", "mean", "median", "p5", "p95", "expected_return", "low_vs_spot",
            "high_vs_spot", "p_below", "p_fall_10", "p_rise_10", "paths"]
    for row, key in zip(range(5, 17), keys, strict=True):
        assert su[f"B{row}"].value == pytest.approx(SIM["scientific"][key], rel=1e-9), key
        assert su[f"C{row}"].value == pytest.approx(SIM["naive"][key], rel=1e-9), key
    assert round(su["B7"].value) == 1300 and su["B13"].value == 0.57


@needs_calc
def test_recalculated_paths_match_the_app(class_book, class_recalculated):
    _, sim, _ = class_book
    _, wb = class_recalculated
    paths = np.array([r for r in wb["Paths"].iter_rows(min_row=2, max_row=62, min_col=2,
                                                        max_col=1001, values_only=True)])
    naive = np.array([r for r in wb["Naive"].iter_rows(min_row=2, max_row=62, min_col=2,
                                                        max_col=1001, values_only=True)])
    assert np.allclose(paths, sim.paths, rtol=1e-12)
    assert np.allclose(naive, sim.naive, rtol=1e-12)


@needs_calc
def test_app_settings_round_trip_through_excel(tmp_path):
    """Live-style data (adjusted closes, Nifty with missing days), EWMA, Euler, 90 days,
    seed 7, news +0.3: the recalculated workbook equals the app."""
    from bmm.data import load_snapshot

    stock = load_snapshot("RELIANCE").adj_close
    index = load_snapshot("^NSEI").adj_close
    a = Assumptions(sigma_method="ewma")
    t, sim, data = build(stock, index, news=0.3, assumptions=a, horizon=90, seed=7,
                         method="Euler")
    path = tmp_path / "app.xlsx"
    path.write_bytes(data)
    report = recalc(path)
    assert report["total_errors"] == 0, report["errors"]
    wb = load_workbook(path, data_only=True)
    wi, su = wb["Inputs"], wb["Summary"]
    assert wi["B17"].value == pytest.approx(t.sigma, rel=1e-10)        # EWMA
    assert wi["B34"].value == pytest.approx(t.beta, rel=1e-10)         # gaps in the Nifty
    assert wi["B28"].value == pytest.approx(t.mu, rel=1e-10)
    s = summary(sim)
    for row, key in ((6, "mean"), (7, "median"), (8, "p5"), (9, "p95"), (13, "p_below")):
        assert su[f"B{row}"].value == pytest.approx(s["scientific"][key], rel=1e-9), key
        assert su[f"C{row}"].value == pytest.approx(s["naive"][key], rel=1e-9), key


@needs_calc
def test_short_history_without_nifty_has_no_errors(tmp_path):
    idx = pd.bdate_range("2026-05-01", periods=100)
    stock = pd.Series(100 * np.exp(np.cumsum(np.random.default_rng(3).normal(0, 0.01, 100))),
                      index=idx)
    t, sim, data = build(stock, None, horizon=20)
    path = tmp_path / "short.xlsx"
    path.write_bytes(data)
    report = recalc(path)
    assert report["total_errors"] == 0, report["errors"]
    wi = load_workbook(path, data_only=True)["Inputs"]
    assert wi["B22"].value == "n/a" and wi["C22"].value == 0      # no 6-month momentum
    assert wi["B34"].value == 1.0                                  # no Nifty: beta 1
    assert wi["B28"].value == pytest.approx(t.mu, rel=1e-10)
